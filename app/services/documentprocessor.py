from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import io
from typing import Dict, Union, Any, Optional
from .dynamictable import generate_dynamic_table

class DocumentProcessor:
    
    def __init__(self):
        pass

    def load_document(self, file_stream: io.BytesIO) -> Document:
        file_stream.seek(0)
        return Document(file_stream)

    def save_document(self, doc: Document) -> io.BytesIO:
        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        return output_stream

    def replace_tags(self, doc: Document, placeholders: Dict[str, Union[str, int, float]]) -> Document:
        nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        
        for sdt in doc.element.findall(".//w:sdt", namespaces=nsmap):
            tag_el = sdt.find(".//w:tag", namespaces=nsmap)
            if tag_el is None:
                continue

            tag_name = tag_el.get("{%s}val" % nsmap["w"])
            if not tag_name:
                continue

            tag_key = tag_name.strip().lower()
            
            matching_key = None
            for placeholder_key in placeholders.keys():
                if placeholder_key.strip().lower() == tag_key:
                    matching_key = placeholder_key
                    break
            
            if matching_key:
                replacement_value = str(placeholders[matching_key])
                
                if '\\n' in replacement_value or '\n' in replacement_value:
                    self._replace_text_with_breaks(sdt, replacement_value, nsmap)
                else:
                    text_el = sdt.find(".//w:t", namespaces=nsmap)
                    if text_el is not None:
                        text_el.text = replacement_value

        return doc
    
    def _replace_text_with_breaks(self, sdt, text: str, nsmap: dict):
        run_el = sdt.find(".//w:r", namespaces=nsmap)
        if run_el is None:
            return
        
        for text_el in run_el.findall(".//w:t", namespaces=nsmap):
            run_el.remove(text_el)
        
        lines = text.replace('\\n', '\n').split('\n')
        
        for i, line in enumerate(lines):
            text_el = OxmlElement('w:t')
            text_el.set(qn('xml:space'), 'preserve')
            text_el.text = line
            run_el.append(text_el)
            
            if i < len(lines) - 1:
                br_el = OxmlElement('w:br')
                run_el.append(br_el)
    
    def process_document(self, document_stream: io.BytesIO, placeholders: Dict[str, Union[str, int, float]], 
                        chart_images: Optional[Dict[str, io.BytesIO]] = None,
                        table_data: Optional[Dict[str, Any]] = None) -> io.BytesIO:
        doc = self.load_document(document_stream)
        doc = self.replace_tags(doc, placeholders)
        
        if table_data:
            tag = table_data.get("tag", "table")
            doc = generate_dynamic_table(doc, tag, table_data)
        
        return self.save_document(doc)
